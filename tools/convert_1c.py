#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Конвертер v2 выгрузки 1С «Остатки и доступность товаров» → data/shop-stock.json

Отличия v2 (требования заказчика):
- ГРУППИРОВКА ПО АРТИКУЛУ: строки 1С «Медаль MMC 1270 (золото/серебро/бронза)»
  собираются в ОДНУ карточку с вариантами — как у поставщика. У каждого
  варианта свой остаток (из 1С) и своё фото (от поставщика).
- ЦЕНЫ: приоритет — публичная цена поставщика (рек. розничная, из снимка
  data/rossimvol-full.json, который готовит tools/parse_rossimvol.py);
  fallback — размерная модель.
- ВИТРИНА: на категорию отбирается PER_CAT артикулов с наибольшим суммарным
  остатком (по умолчанию 20).

Запуск:  python3 tools/convert_1c.py ["Остатки и доступность товара.xlsx"]
"""
import json
import math
import re
import ssl
import sys
import urllib.request
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Нужен openpyxl:  pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "shop-stock.json"
SNAP = ROOT / "data" / "rossimvol-full.json"     # снимок поставщика (фото/цены)
PER_CAT = 20

CYR_TO_LAT = str.maketrans({"А": "A", "В": "B", "С": "C", "Е": "E", "Н": "H",
                            "К": "K", "М": "M", "О": "O", "Р": "P", "Т": "T",
                            "Х": "X", "У": "Y", "Д": "D"})


def norm_art(art: str) -> str:
    return re.sub(r"\s+", " ", art.translate(CYR_TO_LAT)).strip().upper()


def norm_key(art: str) -> str:
    """Канонический ключ = как в снимке поставщика (parse_rossimvol.norm_key)."""
    return re.sub(r"[\s_]+", "", art.upper())

SIZE_PRICE = [
    (re.compile(r"D-?40|04-?40|[\s-]40\b"), 140),
    (re.compile(r"D-?50|-50\b|\b50\b"), 150),
    (re.compile(r"D-?60|-60\b|\b60\b"), 170),
    (re.compile(r"D-?70|-70\b|\b70\b"), 200),
]
DEFAULT_RETAIL = 150

METAL_MAP = {"золот": "gold", "серебр": "silver", "бронз": "bronze", "микс": "mix"}
METAL_SUFFIX = {"gold": "G", "silver": "S", "bronze": "B"}

THEME_MAP = [
    (re.compile(r"VOL|волейбол", re.I), "volleyball"),
    (re.compile(r"FUT|футбол", re.I), "football"),
    (re.compile(r"HOK|хокк", re.I), "hockey"),
    (re.compile(r"ШАХМ|chess", re.I), "chess"),
    (re.compile(r"БЕГ|RUN|марафон|ЛЫЖ|SKI", re.I), "run"),
    (re.compile(r"ПОБЕД|УЧАСТ|ЗА УЧАСТ|АБСОЛ", re.I), "victory"),
    (re.compile(r"ШКОЛ|УЧЕНИ|ЗНАНИЕ", re.I), "school"),
]


def norm_art(art: str) -> str:
    return re.sub(r"\s+", " ", art.translate(CYR_TO_LAT)).strip().upper()


def parse_name(full: str):
    m = re.match(r"^(.*?)\s*\((.*?)\)\s*$", full)
    base, char = (m.group(1), m.group(2)) if m else (full, "")
    art = base.replace("Медаль", "", 1).strip()
    metal = "mix"
    for key, val in METAL_MAP.items():
        if key in char.lower():
            metal = val
            break
    theme = "universal"
    for rx, val in THEME_MAP:
        if rx.search(full):
            theme = val
            break
    return art, metal, char, theme


def fallback_price(art: str) -> int:
    for rx, price in SIZE_PRICE:
        if rx.search(art):
            return price
    return DEFAULT_RETAIL


def load_supplier_snap():
    if not SNAP.exists():
        return {}
    try:
        return json.loads(SNAP.read_text())
    except Exception:
        return {}


UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
_CTX = ssl.create_default_context()


def ensure_local(img_url: str, key: str, suffix: str):
    """Локальный путь к фото варианта; при отсутствии — скачать с поставщика
    (конвертер самодостаточен: работает и до планового прогона парсера)."""
    if not img_url:
        return None
    if not str(img_url).startswith(("http://", "https://", "//")):
        return img_url                       # уже локальный путь
    url = img_url.replace("http://", "https://", 1)
    ext = Path(url.split("?")[0]).suffix or ".jpg"
    name = f"{key.lower()}_{suffix.lower() if suffix else ''}{ext}"
    dst = ROOT / "img" / "remote" / name
    if dst.exists():
        return f"img/remote/{name}"
    try:
        dst.parent.mkdir(parents=True, exist_ok=True)
        req = urllib.request.Request(url, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=30, context=_CTX) as r:
            dst.write_bytes(r.read())
        return f"img/remote/{name}"
    except Exception as e:
        print(f"  ! фото {key}/{suffix}: {e}", file=sys.stderr)
        return None


def supplier_assets(snap, art, metal):
    """Фото и цена варианта от поставщика (по каноническому ключу + металлу)."""
    entry = snap.get(norm_key(art))
    if not entry:
        return None, None
    suffix = METAL_SUFFIX.get(metal)
    price, img = None, None
    for v in entry["variants"]:
        if v.get("suffix") == suffix or (not suffix and not v.get("suffix")):
            price = v.get("supplierPrice")
            img = v.get("img")
            break
    img = ensure_local(img, norm_key(art), suffix) if img else None
    return img, price


def main(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    head_row = None
    for r in range(1, 30):
        if ws.cell(r, 4).value and "Ед" in str(ws.cell(r, 4).value):
            head_row = r
            break
    if not head_row:
        sys.exit("Не нашёл строку заголовков (колонка D «Ед. изм.»)")

    snap = load_supplier_snap()
    groups, warehouse = {}, "Магазин центральный"
    for r in range(head_row + 1, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        name = str(name).strip().rstrip(",")
        unit = ws.cell(r, 4).value
        if name.lower().startswith("итого"):
            continue
        if unit is None:
            warehouse = name if r == head_row + 1 else warehouse
            continue
        avail = int(ws.cell(r, 10).value or 0)
        if avail <= 0:
            continue
        art, metal, char, theme = parse_name(name)
        art = norm_art(art)
        base_name = re.sub(r"\s*\([^)]*\)\s*$", "", name).strip()   # «(бронза)» → вариант
        g = groups.setdefault(art, {
            "art": art, "name": base_name, "cat": "medal", "theme": theme,
            "variants": [],
        })
        g["variants"].append({
            "metal": metal, "char": char,
            "inStock": int(ws.cell(r, 5).value or 0),
            "avail": avail,
        })

    items = []
    n_photo = 0
    for g in groups.values():
        for v in g["variants"]:
            img, price = supplier_assets(snap, g["art"], v["metal"])
            v["img"] = img
            v["price"] = price or fallback_price(g["art"])
            if img:
                n_photo += 1
        g["price"] = min(v["price"] for v in g["variants"])
        g["stock"] = sum(v["avail"] for v in g["variants"])
        items.append(g)

    # витрина: PER_CAT артикулов с наибольшим суммарным остатком в категории
    final = []
    for cat in {i["cat"] for i in items}:
        pool = sorted((i for i in items if i["cat"] == cat),
                      key=lambda x: -x["stock"])
        final += pool[:PER_CAT]

    data = {
        "source": "1С: Остатки и доступность товаров (группировка по артикулам)",
        "warehouse": warehouse,
        "updated": __import__("datetime").datetime.now().isoformat(timespec="seconds"),
        "perCat": PER_CAT,
        "count": len(final),
        "totalArts": len(items),
        "photoCoverage": f"{n_photo}/{sum(len(i['variants']) for i in items)} вариантов с фото поставщика",
        "items": final,
    }
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"OK: {len(items)} артикулов → витрина {len(final)} → {OUT.relative_to(ROOT)}")
    print(f"   фото поставщика подхвачено: {n_photo} вариантов")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "tools" / "1c-stock-sample.xlsx"))
